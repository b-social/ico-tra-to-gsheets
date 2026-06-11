#!/usr/bin/env python3
"""
compare_sheets.py — Compare live sheet vs a freshly-built reference sheet.

Strategy:
  1. Download the live sheet as XLSX (the one the user has modified).
  2. Download a separate reference sheet as XLSX (built by running buildTRATool()
     into a clean test spreadsheet, whose ID you pass as --ref-id).
  3. Diff cell values, merges, and column widths per tab.

Usage:
    # First make sure gcloud auth is current:
    gcloud auth login --enable-gdrive-access

    # Then run (replace REF_ID with the ID of the clean test spreadsheet):
    source .venv/bin/activate
    python compare_sheets.py --ref-id REF_SPREADSHEET_ID

    # To skip re-downloading if you already have the XLSX files:
    python compare_sheets.py --live-xlsx live.xlsx --ref-xlsx reference.xlsx
"""

import argparse
import json
import os
import subprocess
import sys
import urllib.request
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LIVE_SHEET_ID = '1hrlIC5yuqntgBaLRwN57m_uJHin7AY1i0AYrZgkV0Vs'

# Tabs built by the script that we care about (skip user-added tabs like Overview, Q0)
SCRIPT_TABS = [
    '📋 Instructions',
    'Q1 · Transfer Details',
    'Q2 · PI Risk Scores',
    'Q3 · Investigation Level',
    'Q4 · Human Rights Risk',
    'Q5 · Enforcement',
    'Q6 · Exceptions',
    '✅ Summary',
    '🏦 Kroo PI Categories',
    'Lookups',
]

# ── Auth ───────────────────────────────────────────────────────────────────────

def get_token():
    result = subprocess.run(
        ['gcloud', 'auth', 'print-access-token'],
        capture_output=True, text=True
    )
    if result.returncode != 0:
        print("ERROR: gcloud auth failed. Run: gcloud auth login --enable-gdrive-access")
        sys.exit(1)
    return result.stdout.strip()

# ── Download ───────────────────────────────────────────────────────────────────

def download_xlsx(sheet_id, out_path, token):
    """Export a Google Sheet as XLSX using the Drive export endpoint."""
    url = (
        f'https://www.googleapis.com/drive/v3/files/{sheet_id}/export'
        f'?mimeType=application%2Fvnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    req = urllib.request.Request(url, headers={'Authorization': f'Bearer {token}'})
    with urllib.request.urlopen(req) as resp:
        data = resp.read()
    with open(out_path, 'wb') as f:
        f.write(data)
    print(f"  Downloaded {len(data):,} bytes → {out_path}")

# ── Comparison helpers ─────────────────────────────────────────────────────────

def cell_value(cell):
    """Return a normalised string representation of a cell's value."""
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, str) and v.startswith('='):
        return v  # keep formula text as-is for comparison
    return str(v).strip()

def sheet_merges(ws):
    """Return a sorted list of merge range strings."""
    return sorted(str(m) for m in ws.merged_cells.ranges)

def sheet_col_widths(ws):
    """Return dict of {col_letter: width} for columns with explicit widths."""
    widths = {}
    for col_letter, col_dim in ws.column_dimensions.items():
        if col_dim.width:
            widths[col_letter] = round(col_dim.width, 1)
    return widths

def diff_values(ws_live, ws_ref, tab_name, issues):
    """Compare cell values between two worksheets."""
    max_row = max(ws_live.max_row, ws_ref.max_row)
    max_col = max(ws_live.max_column, ws_ref.max_column)
    diffs = 0
    for r in range(1, min(max_row, 200) + 1):
        for c in range(1, min(max_col, 30) + 1):
            live_v = cell_value(ws_live.cell(r, c))
            ref_v  = cell_value(ws_ref.cell(r, c))
            if live_v != ref_v:
                col_l = get_column_letter(c)
                issues.append({
                    'tab': tab_name,
                    'type': 'cell_value',
                    'cell': f'{col_l}{r}',
                    'live': live_v[:120] if len(live_v) > 120 else live_v,
                    'ref':  ref_v[:120]  if len(ref_v)  > 120 else ref_v,
                })
                diffs += 1
                if diffs >= 50:
                    issues.append({'tab': tab_name, 'type': 'truncated',
                                   'note': f'More than 50 value diffs — truncated'})
                    return

def diff_merges(ws_live, ws_ref, tab_name, issues):
    live_m = set(sheet_merges(ws_live))
    ref_m  = set(sheet_merges(ws_ref))
    for m in sorted(live_m - ref_m):
        issues.append({'tab': tab_name, 'type': 'merge_only_in_live', 'range': m})
    for m in sorted(ref_m - live_m):
        issues.append({'tab': tab_name, 'type': 'merge_only_in_ref', 'range': m})

def diff_col_widths(ws_live, ws_ref, tab_name, issues, tolerance=5):
    live_w = sheet_col_widths(ws_live)
    ref_w  = sheet_col_widths(ws_ref)
    all_cols = set(live_w) | set(ref_w)
    for col in sorted(all_cols):
        lw = live_w.get(col, 0)
        rw = ref_w.get(col, 0)
        if abs(lw - rw) > tolerance:
            issues.append({
                'tab': tab_name,
                'type': 'col_width_diff',
                'col': col,
                'live': lw,
                'ref': rw,
            })

# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='Compare live vs reference sheet')
    parser.add_argument('--ref-id',    help='Spreadsheet ID of clean reference build')
    parser.add_argument('--live-xlsx', default='live.xlsx',      help='Path to live XLSX (will download if missing)')
    parser.add_argument('--ref-xlsx',  default='reference.xlsx', help='Path to reference XLSX (will download if missing)')
    parser.add_argument('--no-values', action='store_true', help='Skip cell value diff (faster)')
    args = parser.parse_args()

    token = None

    # Download live sheet
    if not os.path.exists(args.live_xlsx):
        print(f"Downloading live sheet ({LIVE_SHEET_ID}) → {args.live_xlsx}")
        token = token or get_token()
        download_xlsx(LIVE_SHEET_ID, args.live_xlsx, token)
    else:
        print(f"Using cached {args.live_xlsx}")

    # Download reference sheet
    if not os.path.exists(args.ref_xlsx):
        if not args.ref_id:
            print("ERROR: --ref-id required to download reference sheet, or pass --ref-xlsx <path>")
            sys.exit(1)
        print(f"Downloading reference sheet ({args.ref_id}) → {args.ref_xlsx}")
        token = token or get_token()
        download_xlsx(args.ref_id, args.ref_xlsx, token)
    else:
        print(f"Using cached {args.ref_xlsx}")

    print("\nLoading workbooks…")
    wb_live = load_workbook(args.live_xlsx, data_only=False)
    wb_ref  = load_workbook(args.ref_xlsx,  data_only=False)

    issues = []
    summary = {'tabs_in_live': wb_live.sheetnames, 'tabs_in_ref': wb_ref.sheetnames}

    tabs_live = set(wb_live.sheetnames)
    tabs_ref  = set(wb_ref.sheetnames)

    for tab in sorted(tabs_live - tabs_ref):
        issues.append({'tab': tab, 'type': 'tab_only_in_live'})
    for tab in sorted(tabs_ref - tabs_live):
        issues.append({'tab': tab, 'type': 'tab_only_in_ref'})

    common_tabs = [t for t in SCRIPT_TABS if t in tabs_live and t in tabs_ref]
    print(f"Comparing {len(common_tabs)} common script tabs…\n")

    for tab in common_tabs:
        ws_live = wb_live[tab]
        ws_ref  = wb_ref[tab]
        before = len(issues)

        diff_merges(ws_live, ws_ref, tab, issues)
        diff_col_widths(ws_live, ws_ref, tab, issues)
        if not args.no_values:
            diff_values(ws_live, ws_ref, tab, issues)

        n = len(issues) - before
        status = f"{n} issue(s)" if n else "✓ clean"
        print(f"  {tab:<40} {status}")

    # Report
    report_path = 'comparison_report.json'
    with open(report_path, 'w') as f:
        json.dump({'summary': summary, 'issues': issues}, f, indent=2, ensure_ascii=False)

    print(f"\n{'='*60}")
    print(f"Total issues: {len(issues)}")
    print(f"Full report:  {report_path}")

    # Print actionable diffs to stdout
    if issues:
        print("\n── Issues ──")
        for iss in issues:
            t = iss['type']
            tab = iss.get('tab', '')
            if t == 'cell_value':
                print(f"  [{tab}] {iss['cell']}: live={repr(iss['live'][:60])}  ref={repr(iss['ref'][:60])}")
            elif t == 'merge_only_in_live':
                print(f"  [{tab}] MERGE only in live: {iss['range']}")
            elif t == 'merge_only_in_ref':
                print(f"  [{tab}] MERGE only in ref (missing from live): {iss['range']}")
            elif t == 'col_width_diff':
                print(f"  [{tab}] col {iss['col']} width: live={iss['live']}  ref={iss['ref']}")
            elif t in ('tab_only_in_live', 'tab_only_in_ref'):
                print(f"  TAB {t}: {tab}")
            else:
                print(f"  {iss}")

if __name__ == '__main__':
    main()
